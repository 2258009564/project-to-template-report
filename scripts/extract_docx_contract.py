#!/usr/bin/env python3
"""Extract a machine-readable formatting and structure contract from a DOCX."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import zipfile
from collections import Counter, defaultdict
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


CAPTION_RE = re.compile(r"^[\s\u3000]*(?:图|表)\s*[0-9A-Za-z一二三四五六七八九十]+(?:[-－.][0-9A-Za-z一二三四五六七八九十()（）]+)*")


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def value(obj, attr):
    v = getattr(obj, attr, None)
    if v is None:
        return None
    if hasattr(v, "pt"):
        return round(v.pt, 3)
    if hasattr(v, "emu"):
        return int(v.emu)
    if hasattr(v, "rgb"):
        return str(v.rgb) if v.rgb is not None else None
    try:
        return int(v)
    except (TypeError, ValueError):
        return str(v)


def clean_text(text: str, limit: int = 120) -> str:
    text = re.sub(r"\s+", " ", text).strip()
    return text[:limit] + ("..." if len(text) > limit else "")


def font_contract(font) -> dict:
    element = getattr(font, "_element", None)
    rpr = getattr(element, "rPr", None) if element is not None else None
    rfonts = rpr.find(qn("w:rFonts")) if rpr is not None else None
    return {
        "name": font.name,
        "ascii": rfonts.get(qn("w:ascii")) if rfonts is not None else None,
        "hAnsi": rfonts.get(qn("w:hAnsi")) if rfonts is not None else None,
        "eastAsia": rfonts.get(qn("w:eastAsia")) if rfonts is not None else None,
        "size_pt": value(font, "size"),
        "bold": font.bold,
        "italic": font.italic,
        "underline": str(font.underline) if font.underline is not None else None,
        "color": value(font.color, "rgb") if getattr(font, "color", None) else None,
    }


def paragraph_format_contract(pf) -> dict:
    return {
        "alignment": value(pf, "alignment"),
        "left_indent_pt": value(pf, "left_indent"),
        "right_indent_pt": value(pf, "right_indent"),
        "first_line_indent_pt": value(pf, "first_line_indent"),
        "space_before_pt": value(pf, "space_before"),
        "space_after_pt": value(pf, "space_after"),
        "line_spacing": value(pf, "line_spacing"),
        "line_spacing_rule": value(pf, "line_spacing_rule"),
        "keep_together": pf.keep_together,
        "keep_with_next": pf.keep_with_next,
        "page_break_before": pf.page_break_before,
        "widow_control": pf.widow_control,
    }


def cell_shading(cell) -> str | None:
    shd = cell._tc.tcPr.find(qn("w:shd"))
    return shd.get(qn("w:fill")) if shd is not None else None


def border_summary(tbl) -> dict:
    borders = tbl._tbl.tblPr.find(qn("w:tblBorders"))
    result = {}
    if borders is None:
        return result
    for child in borders:
        name = child.tag.rsplit("}", 1)[-1]
        result[name] = {
            "val": child.get(qn("w:val")),
            "size": child.get(qn("w:sz")),
            "color": child.get(qn("w:color")),
        }
    return result


def cached_page_count(path: Path) -> int | None:
    try:
        with zipfile.ZipFile(path) as zf:
            data = zf.read("docProps/app.xml").decode("utf-8", "ignore")
        match = re.search(r"<Pages>(\d+)</Pages>", data)
        return int(match.group(1)) if match else None
    except (KeyError, zipfile.BadZipFile):
        return None


def extract(path: Path, max_samples: int) -> dict:
    doc = Document(path)
    style_counts = Counter()
    style_samples: dict[str, list[str]] = defaultdict(list)
    headings = []
    captions = []
    paragraph_samples = {}

    for index, p in enumerate(doc.paragraphs, 1):
        style = p.style.name if p.style is not None else ""
        text = clean_text(p.text)
        style_counts[style] += 1
        if text and len(style_samples[style]) < max_samples:
            style_samples[style].append(text)
            paragraph_samples.setdefault(style, []).append({
                "paragraph": index, "text": text,
                "direct_paragraph_format": paragraph_format_contract(p.paragraph_format),
                "runs": [{"text": clean_text(r.text, 60), "direct_font": font_contract(r.font)}
                         for r in p.runs if r.text.strip()][:8],
            })
        if style.lower().startswith("heading") or style.startswith("标题"):
            headings.append({"paragraph": index, "style": style, "text": text})
        if style.lower() == "caption" or style == "题注" or CAPTION_RE.match(text):
            captions.append({"paragraph": index, "style": style, "text": text})

    styles = {}
    for st in doc.styles:
        if st.type != 1 and st.name not in style_counts:
            continue
        styles[st.name] = {
            "type": int(st.type),
            "used_count": style_counts.get(st.name, 0),
            "base_style": st.base_style.name if st.base_style else None,
            "next_style": st.next_paragraph_style.name if getattr(st, "next_paragraph_style", None) else None,
            "font": font_contract(st.font),
            "paragraph_format": paragraph_format_contract(st.paragraph_format) if st.type == 1 else None,
            "samples": style_samples.get(st.name, []),
        }

    sections = []
    for sec in doc.sections:
        sections.append({
            "orientation": value(sec, "orientation"),
            "page_width_pt": value(sec, "page_width"),
            "page_height_pt": value(sec, "page_height"),
            "top_margin_pt": value(sec, "top_margin"),
            "bottom_margin_pt": value(sec, "bottom_margin"),
            "left_margin_pt": value(sec, "left_margin"),
            "right_margin_pt": value(sec, "right_margin"),
            "header_distance_pt": value(sec, "header_distance"),
            "footer_distance_pt": value(sec, "footer_distance"),
            "start_type": value(sec, "start_type"),
            "different_first_page": sec.different_first_page_header_footer,
            "header_text": clean_text(" ".join(p.text for p in sec.header.paragraphs), 500),
            "footer_text": clean_text(" ".join(p.text for p in sec.footer.paragraphs), 500),
        })

    tables = []
    shading = Counter()
    for i, tbl in enumerate(doc.tables, 1):
        for row in tbl.rows:
            for cell in row.cells:
                shading[cell_shading(cell) or "none"] += 1
        tables.append({
            "index": i,
            "rows": len(tbl.rows),
            "columns": len(tbl.columns),
            "style": tbl.style.name if tbl.style else None,
            "alignment": value(tbl, "alignment"),
            "autofit": tbl.autofit,
            "borders": border_summary(tbl),
            "header_cells": [clean_text(c.text, 80) for c in tbl.rows[0].cells] if tbl.rows else [],
        })

    root = doc._element
    fields = []
    for instr in root.xpath('.//*[local-name()="instrText"]'):
        if instr.text:
            fields.append(clean_text(instr.text, 160))

    return {
        "schema_version": 1,
        "source": {
            "path": str(path.resolve()),
            "bytes": path.stat().st_size,
            "sha256": sha256(path),
            "cached_pages": cached_page_count(path),
        },
        "counts": {
            "sections": len(doc.sections),
            "paragraphs": len(doc.paragraphs),
            "tables": len(doc.tables),
            "inline_images": len(doc.inline_shapes),
            "floating_images": len(root.xpath('.//*[local-name()="anchor"]')),
            "bookmarks": len(root.xpath('.//*[local-name()="bookmarkStart"]')),
            "content_controls": len(root.xpath('.//*[local-name()="sdt"]')),
            "fields": len(fields),
        },
        "sections": sections,
        "styles": styles,
        "direct_format_samples": paragraph_samples,
        "format_resolution": "OOXML direct values; null means inherited, not absent. Use Word COM profile for effective values.",
        "style_counts": dict(style_counts),
        "headings": headings,
        "captions": captions,
        "tables": tables,
        "cell_shading_counts": dict(shading),
        "field_instructions": fields,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("--out", type=Path)
    parser.add_argument("--max-samples", type=int, default=3)
    args = parser.parse_args()
    if not args.input.is_file():
        raise SystemExit(f"input not found: {args.input}")
    data = extract(args.input, args.max_samples)
    output = args.out or args.input.with_suffix(".contract.json")
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    print(output.resolve())


if __name__ == "__main__":
    main()
